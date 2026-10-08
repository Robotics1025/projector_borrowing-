from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "coursework_slides.pptx"

NAVY = RGBColor(15, 31, 56)
BLUE = RGBColor(43, 108, 176)
CYAN = RGBColor(49, 196, 190)
WHITE = RGBColor(247, 250, 252)
MUTED = RGBColor(174, 190, 209)


SLIDES = [
    ("Projector Borrowing", ["DDD · TDD · Clean Architecture", "Makerere University — Group Coursework", "Member 1 — [Name] · Domain Lead", "Member 2 — [Name] · Domain Model Lead", "Member 3 — [Name] · Architecture Lead", "Member 4 — [Name] · Testing Lead", "Member 5 — [Name] · Integration Lead"]),
    ("Problem & Scope", ["Problem — manual projector lending can allow clashes, invalid durations, and unclear accountability.", "Inside — request and issue a loan; return is retained as a supporting workflow; in-memory persistence.", "Outside — authentication, payments, web UI, production database, deployment, messaging.", "Terms — Borrower, Loan, Borrowing Period, Projector, Loan Status, Projector Status."]),
    ("Six Business Rules", ["BR1 · BorrowingPeriod: 1–7 calendar days; otherwise InvalidBorrowingPeriodError.", "BR2 · Projector: AVAILABLE → BORROWED only; otherwise InvalidProjectorTransitionError.", "BR3 · Borrower: at most two PENDING/ACTIVE loans; otherwise LoanLimitExceededError.", "BR4 · Eligibility Service: students cannot borrow PREMIUM; otherwise PremiumProjectorRestrictedError.", "BR5 · Borrower raises LoanIssued; handler asks Projector to mark itself BORROWED; rejection propagates.", "BR6 · RequestLoan loads both aggregates first; missing item produces Borrower/ProjectorNotFoundError."]),
    ("Value Object & Entity", ["BorrowingPeriod enforces BR1 at construction and is immutable.", "It has no identity: equality depends only on start and due dates.", "Loan is an Entity: its UUID preserves identity while status changes.", "Projector is an Entity and Aggregate Root: projector_id defines identity; BR2 protects its lifecycle."]),
    ("Two Aggregates", ["Aggregate A — Borrower root contains its Loan entities.", "Invariant (BR3) — no more than two open loans; Loan IDs are unique within the borrower.", "Aggregate B — Projector root owns type, availability status, and version.", "Invariant (BR2) — only a currently available projector may become borrowed.", "No aggregate directly mutates the internals of the other."]),
    ("Domain Service & Factory", ["BorrowingEligibilityService owns BR4 because the decision needs Borrower.role and Projector.type.", "It also confirms current availability and capacity at the decision point.", "The service is stateless and contains domain policy only.", "Factory rejected: creation is simple; constructors and aggregate methods already enforce all creation rules."]),
    ("Repositories", ["BorrowerRepository — get_by_id and save; stores Borrower aggregates.", "ProjectorRepository — get_by_id and version-aware save; stores Projector aggregates.", "InMemoryBorrowerRepository and InMemoryProjectorRepository are infrastructure adapters.", "Deep copies simulate persistence boundaries; projector saves reject stale versions."]),
    ("Clean Architecture", ["INTERFACE  →  APPLICATION  →  DOMAIN", "     │                 ↑", "     └─ composition ─ INFRASTRUCTURE", "Domain — rules, aggregates, value object, service, events, repository protocols.", "Application — use cases, event handler, input/output DTOs.", "Infrastructure — in-memory repository implementations. Interface — console presenter/entry point.", "Layer Supertype rejected: no shared domain behaviour justifies a base class."]),
    ("Main Application Service", ["IssueLoan coordinates: load → eligibility check → issue → dispatch event → save.", "Input DTO — IssueLoanCommand(borrower_id, loan_id).", "Output DTO — LoanResult(loan_id, projector_id, status, dates).", "Dependency injection — repository implementations enter through the constructor.", "Core rules remain in Borrower, Projector, BorrowingPeriod, and the domain service."]),
    ("BR5 Domain Event Flow", ["IssueLoanCommand → IssueLoan → Borrower.issue_loan", "Borrower changes Loan PENDING → ACTIVE, then raises LoanIssued.", "LoanIssuedHandler loads the Projector and calls mark_borrowed().", "Projector checks BR2 before accepting AVAILABLE → BORROWED.", "The flow is synchronous and in-process; the aggregates remain decoupled."]),
    ("T1–T8 Test Map", ["T1 BR1 — accepts 7-day boundary; rejects 8 days.  T2 BR2 — rejects repeated borrow transition.", "T3 BR3 — rejects third open loan.  T4 BR4 — rejects student + premium.", "T5 BR5 — verifies LoanIssued payload.  T6 BR6 — rejects missing borrower lookup.", "T7 — issuing handles the event and stores Projector as BORROWED.", "T8 — Projector rejects follow-up; state/version remain unchanged.", "Expected result: all eight pass; rejection and boundary requirements are covered."]),
    ("TDD Example — BR1", ["RED — T1 expected the 7-day boundary to be valid and 8 days to be rejected.", "Implementation — BorrowingPeriod validates: 1 <= duration_in_days() <= 7.", "GREEN — T1 passes with the invariant located in the Value Object.", "Actual red/green terminal output is kept in evidence/tdd_cycle.txt."]),
    ("Main Use Case Walkthrough", ["1 · Interface/composition root builds IssueLoan with both in-memory repositories.", "2 · IssueLoanCommand identifies the borrower and pending loan.", "3 · Repositories load Borrower and Projector (BR6).", "4 · Eligibility service evaluates cross-concept policy (BR4).", "5 · Borrower activates Loan and records LoanIssued (BR5).", "6 · Handler asks Projector to become BORROWED (BR2), repositories save, LoanResult returns."]),
    ("Design Choice Changed", ["Rejected design — IssueLoan directly called projector.mark_borrowed() after changing Borrower.", "Problem — the event was only recorded; it did not drive the follow-up action required by BR5.", "Change — LoanIssuedHandler now consumes LoanIssued and coordinates Aggregate B.", "Benefit — explicit, testable event flow without direct aggregate-to-aggregate mutation.", "Trade-off — synchronous in-memory dispatch is intentionally small for coursework scope."]),
    ("Traceability & Conclusion", ["BR1 → BorrowingPeriod → T1 → pass.  BR2 → Projector → T2/T8 → pass.", "BR3 → Borrower → T3 → pass.  BR4 → Eligibility Service → T4 → pass.", "BR5 → LoanIssued + handler → T5/T7 → pass.  BR6 → repositories + RequestLoan → T6 → pass.", "Two aggregates, inward dependencies, injected repositories, DTOs, and in-process event handling match the implementation.", "Conclusion — small scope, rules in the domain, application coordination only, automated evidence retained."]),
]


def add_slide(prs: Presentation, number: int, title: str, bullets: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = NAVY

    accent = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(0.18), Inches(7.5))
    accent.fill.solid()
    accent.fill.fore_color.rgb = CYAN
    accent.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.75), Inches(0.55), Inches(11.7), Inches(0.8))
    title_frame = title_box.text_frame
    title_frame.text = title
    title_frame.paragraphs[0].font.name = "Aptos Display"
    title_frame.paragraphs[0].font.size = Pt(30)
    title_frame.paragraphs[0].font.bold = True
    title_frame.paragraphs[0].font.color.rgb = WHITE

    body = slide.shapes.add_textbox(Inches(0.9), Inches(1.55), Inches(11.5), Inches(5.25))
    frame = body.text_frame
    frame.word_wrap = True
    for index, bullet in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = bullet
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(18 if len(bullets) <= 5 else 16)
        paragraph.font.color.rgb = WHITE
        paragraph.space_after = Pt(10)
        paragraph.level = 0

    footer = slide.shapes.add_textbox(Inches(0.85), Inches(7.05), Inches(11.5), Inches(0.25))
    footer_frame = footer.text_frame
    footer_frame.text = f"PROJECTOR BORROWING  /  {number:02d}"
    footer_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    footer_frame.paragraphs[0].font.size = Pt(9)
    footer_frame.paragraphs[0].font.color.rgb = MUTED


def main() -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    for number, (title, bullets) in enumerate(SLIDES, start=1):
        add_slide(presentation, number, title, bullets)
    presentation.save(OUTPUT)


if __name__ == "__main__":
    main()
